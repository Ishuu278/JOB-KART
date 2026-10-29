# -*- coding: utf-8 -*-
"""JOB-KART academic presentation generator (16:9, 10 slides)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

ASSETS = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "ppt-assets"))
CROPS = os.path.join(ASSETS, "crops")
OUT = os.environ.get("PPT_OUT") or os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "JOB-KART-Presentation.pptx"))
os.makedirs(CROPS, exist_ok=True)

# ---------- wide crops (browser-style band from each screenshot) ----------
def crop(src, dst, anchor="top", y0=0):
    im = Image.open(os.path.join(ASSETS, src))
    w, h = im.size
    th = int(w / 2.6)  # target aspect 2.6:1
    if anchor == "top":
        box = (0, 0, w, min(th, h))
    else:
        y0 = min(y0, h - th)
        box = (0, max(0, y0), w, max(0, y0) + th)
    im.crop(box).save(os.path.join(CROPS, dst))

crop("18-candidate-dash-v2.png",    "s3-candidate.png")
crop("19-recruiter-dash-v4.png",      "s3-recruiter.png")
crop("06-job-details.png",           "s4-details.png")
crop("12-company.png",               "s4-company.png")
crop("09-candidate-dash-charts.png", "s4-ccharts.png", anchor="band", y0=170)
crop("11-recruiter-charts.png",      "s4-rcharts.png", anchor="band", y0=150)
crop("20-ats-scanner.png",            "s6-ats.png", anchor="band", y0=78)

# ---------- palette ----------
DARK   = RGBColor(0x0A, 0x1F, 0x44)
DARK2  = RGBColor(0x11, 0x2B, 0x55)
DARK3  = RGBColor(0x16, 0x35, 0x66)
BLUE   = RGBColor(0x25, 0x63, 0xEB)
LBLUE  = RGBColor(0x5B, 0x9B, 0xF5)
SKY    = RGBColor(0xEA, 0xF1, 0xFE)
PAPER  = RGBColor(0xF6, 0xF8, 0xFC)
INK    = RGBColor(0x0F, 0x1D, 0x33)
MUTED  = RGBColor(0x5A, 0x6B, 0x87)
AMBER  = RGBColor(0xF5, 0x9E, 0x0B)
AMBERD = RGBColor(0xB4, 0x74, 0x06)
GREEN  = RGBColor(0x10, 0xB9, 0x81)
PURPLE = RGBColor(0x7C, 0x3A, 0xED)
TEAL   = RGBColor(0x0E, 0x74, 0x9C)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LINEC  = RGBColor(0xDD, 0xE5, 0xF2)
FONT   = "Segoe UI"

TEAM = [
    ("Tusar Kumar Mahakud",    "2301289131", "TM"),
    ("Swadhin Panigrahi",      "2301289125", "SP"),
    ("Tannu Kumari",           "2301289130", "TK"),
    ("Subhrakant Senapati",    "2301289122", "SS"),
]

SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]

# ---------- helpers ----------
def slide():
    return prs.slides.add_slide(BLANK)

def set_alpha(shape, pct):
    sF = shape.fill._xPr.find(qn('a:solidFill'))
    clr = sF.find(qn('a:srgbClr'))
    a = clr.makeelement(qn('a:alpha'), {'val': str(int(pct * 1000))})
    clr.append(a)

def rect(s, x, y, w, h, fill, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None, alpha=None):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if alpha is not None:
        set_alpha(sp, alpha)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp

def circle(s, cx, cy, r, fill, alpha=None, line=None, lw=1.0):
    sp = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - r, cy - r, r * 2, r * 2)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if alpha is not None:
        set_alpha(sp, alpha)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp

def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf

def para(tf, first=False, align=PP_ALIGN.LEFT, space_before=0, space_after=0, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    return p

def run(p, text, size, color, bold=False, italic=False, spacing=None, font=FONT):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    if spacing is not None:
        r.font._rPr.set('spc', str(int(spacing * 100)))
    return r

def kicker(s, x, y, num, label, color=BLUE, on_dark=False):
    t = tb(s, x, y, Inches(9), Inches(0.3))
    p = para(t, first=True)
    run(p, num + "  ", 12, color, bold=True, spacing=2)
    run(p, "—  " + label.upper(), 12, LBLUE if on_dark else MUTED, bold=True, spacing=2)

def title(s, x, y, w, text, size=30, color=INK, h=0.8):
    t = tb(s, x, y, w, Inches(h))
    p = para(t, first=True, line=1.02)
    run(p, text, size, color, bold=True)

def subtitle(s, x, y, w, text, color=MUTED, size=12.5):
    t = tb(s, x, y, w, Inches(0.55))
    p = para(t, first=True, line=1.2)
    run(p, text, size, color, italic=True)

def pagenum(s, n, on_dark=False):
    t = tb(s, SW - Inches(1.5), SH - Inches(0.45), Inches(1.1), Inches(0.28))
    p = para(t, first=True, align=PP_ALIGN.RIGHT)
    run(p, f"{n:02d} / 10", 10, LBLUE if on_dark else MUTED, bold=True, spacing=1)

def brandfoot(s, on_dark=False):
    t = tb(s, Inches(0.55), SH - Inches(0.45), Inches(6), Inches(0.28))
    p = para(t, first=True)
    run(p, "JOB", 10, WHITE if on_dark else INK, bold=True, spacing=3)
    run(p, "·", 10, LBLUE, bold=True)
    run(p, "KART", 10, WHITE if on_dark else INK, bold=True, spacing=3)
    run(p, "   |   Empowering Careers, Simplifying Hiring", 9, LBLUE if on_dark else MUTED, italic=True)

BAR_H = Inches(0.30)
def browser_frame(s, img, x, y, w, aspect=2.6):
    """Browser-chrome framed screenshot; returns bottom y."""
    img_h = Emu(int(w / aspect))
    card = rect(s, x - Inches(0.05), y - Inches(0.05), w + Inches(0.10),
                Emu(int(img_h) + int(BAR_H) + Inches(0.10)),
                WHITE, line=LINEC, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(s, x, y, w, BAR_H, RGBColor(0xE9, 0xEE, 0xF7), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    for i, c in enumerate((RGBColor(0xFF, 0x5F, 0x57), RGBColor(0xFE, 0xBC, 0x2E), RGBColor(0x28, 0xC8, 0x40))):
        circle(s, x + Inches(0.22 + i * 0.19), y + BAR_H / 2, Inches(0.042), c)
    rect(s, x + Inches(0.8), y + Inches(0.065), w - Inches(1.05), BAR_H - Inches(0.13),
         WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    t = tb(s, x + Inches(0.9), y + Inches(0.055), w - Inches(1.2), BAR_H - Inches(0.1))
    p = para(t, first=True)
    run(p, "job-kart.org", 8, MUTED)
    s.shapes.add_picture(img, x, y + BAR_H, width=w, height=img_h)
    return Emu(int(y) + int(BAR_H) + int(img_h))

def chip(s, x, y, text, fill, color, size=10.5, bold=True, h=0.34, pad=0.22):
    w = Inches(pad * 2 + len(text) * 0.082 * (size / 11.0))
    sp = rect(s, x, y, w, Inches(h), fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = sp.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, text, size, color, bold=bold)
    return x + w + Inches(0.14)

def bullet(tf, head, body, mark="●", mark_color=BLUE, size=12, dark=False, first=False, gap=7):
    p = para(tf, first=first, space_after=gap, line=1.1)
    run(p, mark + "  ", size - 3, mark_color, bold=True)
    run(p, head + "  ", size, WHITE if dark else INK, bold=True)
    if body:
        run(p, body, size - 1.5, LBLUE if dark else MUTED)
    return p

def arrow(s, x, y, w, color, h=0.026):
    ln = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y - Inches(h), w, Inches(h * 2))
    ln.fill.solid(); ln.fill.fore_color.rgb = color
    ln.line.fill.background(); ln.shadow.inherit = False
    try:
        ln.adjustments[0] = 0.62; ln.adjustments[1] = 0.55
    except Exception:
        pass
    return ln

C = lambda name: os.path.join(CROPS, name)
A = lambda name: os.path.join(ASSETS, name)

# ============================================================
# SLIDE 1 — TITLE (project, team names & registration numbers)
# ============================================================
s = slide()
rect(s, 0, 0, SW, SH, DARK)
circle(s, SW - Inches(1.2), Inches(-0.6), Inches(2.6), DARK3, alpha=55)
circle(s, Inches(-0.9), SH + Inches(0.4), Inches(2.4), DARK3, alpha=45)
circle(s, SW - Inches(2.1), SH + Inches(0.7), Inches(1.5), BLUE, alpha=18)

t = tb(s, Inches(0.62), Inches(0.62), Inches(6.6), Inches(0.4))
p = para(t, first=True)
run(p, "FULL-STACK PROJECT PRESENTATION", 12, LBLUE, bold=True, spacing=3)

t = tb(s, Inches(0.6), Inches(1.05), Inches(6.7), Inches(1.4))
p = para(t, first=True, line=1.0)
run(p, "JOB", 58, WHITE, bold=True)
run(p, "·", 58, LBLUE, bold=True)
run(p, "KART", 58, WHITE, bold=True)

t = tb(s, Inches(0.63), Inches(2.22), Inches(6.4), Inches(0.45))
p = para(t, first=True)
run(p, "Empowering Careers, Simplifying Hiring", 16, LBLUE, bold=True, italic=True)

t = tb(s, Inches(0.63), Inches(2.72), Inches(6.3), Inches(0.85))
p = para(t, first=True, line=1.25)
run(p, "A full-stack job-portal platform connecting verified talent with India's "
       "top employers — built on the MERN-style stack.", 12.5, RGBColor(0xC9, 0xD6, 0xEE))

x, y = Inches(0.63), Inches(3.72)
for txt in ("MongoDB", "Express.js 5", "Socket.io", "JWT Auth", "Vanilla JS"):
    x = chip(s, x, y, txt, DARK3, RGBColor(0xBF, 0xD3, 0xF5), h=0.34)

# team cards
t = tb(s, Inches(0.63), Inches(4.12), Inches(5), Inches(0.32))
p = para(t, first=True)
run(p, "PRESENTED BY", 11, RGBColor(0x8F, 0xA6, 0xCB), bold=True, spacing=3)

tcw, tch = Inches(3.42), Inches(1.02)
for i, (name, reg, init) in enumerate(TEAM):
    xx = Inches(0.6) + (i % 2) * (tcw + Inches(0.18))
    yy = Inches(4.48) + (i // 2) * (tch + Inches(0.16))
    rect(s, xx, yy, tcw, tch, DARK2, line=RGBColor(0x2E, 0x4C, 0x85), lw=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    circle(s, xx + Inches(0.42), yy + tch / 2, Inches(0.26), BLUE)
    t = tb(s, xx + Inches(0.16), yy + tch / 2 - Inches(0.13), Inches(0.52), Inches(0.3))
    p = para(t, first=True, align=PP_ALIGN.CENTER)
    run(p, init, 12, WHITE, bold=True)
    t = tb(s, xx + Inches(0.78), yy + Inches(0.18), tcw - Inches(0.9), Inches(0.45))
    p = para(t, first=True, line=1.0)
    run(p, name, 11.5, WHITE, bold=True)
    t = tb(s, xx + Inches(0.78), yy + Inches(0.6), tcw - Inches(0.9), Inches(0.3))
    p = para(t, first=True)
    run(p, "Regd No. " + reg, 9.5, LBLUE, bold=True, spacing=1)

t = tb(s, Inches(0.63), SH - Inches(0.5), Inches(6), Inches(0.3))
p = para(t, first=True)
run(p, "Guide: —   ·   Session 2024–26   ·   2026", 10, RGBColor(0x8F, 0xA6, 0xCB), bold=True, spacing=1)

# right collage: hero browser + phone below
browser_frame(s, A("01-home-hero.png"), Inches(8.05), Inches(0.7), Inches(4.7), aspect=16 / 10)
ph_w = Inches(1.2)
ph_h = Emu(int(int(ph_w) * 844 / 390))
ph_x, ph_y = Inches(11.55), Inches(4.15)
rect(s, ph_x - Inches(0.07), ph_y - Inches(0.07), ph_w + Inches(0.14), ph_h + Inches(0.14),
     DARK3, line=RGBColor(0x2E, 0x4C, 0x85), lw=1.2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
s.shapes.add_picture(A("13-mobile-home.png"), ph_x, ph_y, width=ph_w, height=ph_h)
t = tb(s, Inches(8.9), Inches(5.2), Inches(2.55), Inches(0.3))
p = para(t, first=True, align=PP_ALIGN.RIGHT)
run(p, "Mobile-first UX  \u2192", 9.5, RGBColor(0x9F, 0xB4, 0xD8), italic=True)

pagenum(s, 1, on_dark=True)

# ============================================================
# SLIDE 2 — PROBLEM STATEMENT & MOTIVATION
# ============================================================
s = slide()
rect(s, 0, 0, SW, SH, PAPER)
rect(s, 0, 0, Inches(0.14), SH, BLUE)
kicker(s, Inches(0.6), Inches(0.42), "01", "Problem Statement & Motivation")
title(s, Inches(0.6), Inches(0.72), Inches(12.1), "Hiring is fragmented, opaque and slow", size=26)
subtitle(s, Inches(0.6), Inches(1.2), Inches(7.5),
         "Motivation: students chase scattered postings while recruiters juggle spreadsheets — "
         "JOB·KART was built to fix both sides of this broken loop.")

pains = [
    ("Scattered Listings", "Openings are spread across dozens of boards and groups — no single, verified source of truth for job seekers.", "⌖", BLUE),
    ("Zero Transparency", "Candidates apply into a void: no status, no timeline, no feedback loop after clicking apply.", "?", LBLUE),
    ("Manual Screening", "Recruiters drown in unstructured CVs and spreadsheets, tracking who is applied, shortlisted or hired.", "≡", AMBER),
    ("Employer Trust Gap", "Postings reveal little about companies — no profiles, culture or verified employee reviews to rely on.", "⚑", GREEN),
]
cw, ch = Inches(3.62), Inches(2.22)
gx, gy = Inches(0.6), Inches(1.95)
for i, (h1, b, ic, c) in enumerate(pains):
    xx = gx + (i % 2) * (cw + Inches(0.3))
    yy = gy + (i // 2) * (ch + Inches(0.26))
    rect(s, xx, yy, cw, ch, WHITE, line=LINEC, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    rect(s, xx, yy, Inches(0.09), ch, c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    t = tb(s, xx + Inches(0.28), yy + Inches(0.18), cw - Inches(0.5), Inches(0.45))
    p = para(t, first=True)
    run(p, ic + "  ", 15, c, bold=True)
    run(p, h1, 14.5, INK, bold=True)
    t2 = tb(s, xx + Inches(0.28), yy + Inches(0.66), cw - Inches(0.52), ch - Inches(0.8))
    p2 = para(t2, first=True, line=1.18)
    run(p2, b, 10.8, MUTED)

bot = browser_frame(s, A("04-openings.png"), Inches(8.42), Inches(1.32), Inches(4.45), aspect=16 / 10)
rect(s, Inches(8.1), bot + Inches(0.26), Inches(5.05), Inches(0.6), DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.35)
t = tb(s, Inches(8.3), bot + Inches(0.26), Inches(4.65), Inches(0.6), anchor=MSO_ANCHOR.MIDDLE)
p = para(t, first=True, align=PP_ALIGN.CENTER)
run(p, "18 live roles · 7 companies · 0 status updates after applying", 11, WHITE, bold=True)

sy3 = bot + Inches(1.12)
rect(s, Inches(8.1), sy3, Inches(5.05), Inches(1.72), SKY, line=RGBColor(0xC7, 0xDA, 0xF8), lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
t = tb(s, Inches(8.38), sy3 + Inches(0.16), Inches(4.5), Inches(0.3))
p = para(t, first=True)
run(p, "WHY IT MATTERS", 11, RGBColor(0x1D, 0x4E, 0xD8), bold=True, spacing=2)
t = tb(s, Inches(8.38), sy3 + Inches(0.5), Inches(4.5), Inches(1.1))
p = para(t, first=True, line=1.25)
run(p, "India's hiring market serves millions of job-seekers every year — every missing "
       "status update and every manual spreadsheet multiplies wasted effort on both sides "
       "of the desk.", 11, RGBColor(0x2B, 0x3E, 0x5E))

brandfoot(s)
pagenum(s, 2)

# ============================================================
# SLIDE 3 — OBJECTIVES
# ============================================================
s = slide()
rect(s, 0, 0, SW, SH, WHITE)
rect(s, SW - Inches(0.14), 0, Inches(0.14), SH, BLUE)
kicker(s, Inches(0.6), Inches(0.42), "02", "Objectives")
title(s, Inches(0.6), Inches(0.72), Inches(9.5), "What the project set out to achieve")

objectives = [
    ("Unified verified marketplace", "One trusted platform aggregating verified openings from real companies.", BLUE),
    ("Transparent application tracking", "A real-time status pipeline — Applied → Under Review → Shortlisted → Interview → Selected — for every candidate.", LBLUE),
    ("Efficient hiring workflow", "Visual pipeline and screening tools that replace manual spreadsheet tracking for recruiters.", AMBER),
    ("Insightful dashboards", "Live analytics charts for applications, hiring status and per-job performance.", GREEN),
    ("Secure role-based access", "JWT authentication with candidate/recruiter separation and hashed credentials.", PURPLE),
    ("Lightweight responsive UX", "Framework-free, mobile-first interface with dark mode and zero build step.", TEAL),
]
ocw, och = Inches(6.02), Inches(1.5)
oy = Inches(1.72)
for i, (h1, b, c) in enumerate(objectives):
    xx = Inches(0.6) + (i % 2) * (ocw + Inches(0.28))
    yy = oy + (i // 2) * (och + Inches(0.22))
    rect(s, xx, yy, ocw, och, PAPER, line=LINEC, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
    t = tb(s, xx + Inches(0.26), yy + Inches(0.24), Inches(0.85), Inches(0.7))
    p = para(t, first=True)
    run(p, f"{i+1:02d}", 26, c, bold=True)
    t = tb(s, xx + Inches(1.05), yy + Inches(0.2), ocw - Inches(1.3), Inches(0.4))
    p = para(t, first=True)
    run(p, h1, 13.5, INK, bold=True)
    t = tb(s, xx + Inches(1.05), yy + Inches(0.62), ocw - Inches(1.3), och - Inches(0.75))
    p = para(t, first=True, line=1.15)
    run(p, b, 10.8, MUTED)

t = tb(s, Inches(0.6), Inches(6.72), Inches(12), Inches(0.4))
p = para(t, first=True)
run(p, "Scope: ", 11, INK, bold=True)
run(p, "full-stack web app — REST API, authenticated realtime, ATS resume intelligence + responsive frontend; self-hostable on any Node.js runtime with WebSocket support.", 11, MUTED)

pagenum(s, 3)

# ============================================================
# SLIDE 4 — FUNCTIONALITIES · CANDIDATES
# ============================================================
s = slide()
rect(s, 0, 0, SW, SH, WHITE)
rect(s, 0, 0, SW, Inches(0.12), BLUE)
kicker(s, Inches(0.6), Inches(0.4), "03", "Project Functionalities · Candidates")
title(s, Inches(0.6), Inches(0.7), Inches(12.2), "For candidates — discover, apply, track and get hired", size=26)

cy, ch2 = Inches(1.5), Inches(4.75)
cx, cw2 = Inches(0.6), Inches(6.05)
rect(s, cx, cy, cw2, ch2, SKY, line=RGBColor(0xC7, 0xDA, 0xF8), lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.055)
t = tb(s, cx + Inches(0.3), cy + Inches(0.2), cw2 - Inches(0.6), Inches(0.35))
p = para(t, first=True)
run(p, "FOR CANDIDATES", 12.5, BLUE, bold=True, spacing=2)
tf = tb(s, cx + Inches(0.3), cy + Inches(0.66), cw2 - Inches(0.55), ch2 - Inches(0.8))
bullet(tf, "Smart discovery", "— keyword search, type, location & fresher filters, salary sort", first=True, size=11, gap=7)
bullet(tf, "One-click apply", "— auto-filled profile + PDF resume upload", size=11, gap=7)
bullet(tf, "Deadline-aware apply", "— expired roles hidden, apply window enforced", size=11, gap=7)
bullet(tf, "ATS scanner", "— resume score, fixes & job-fit ranking", size=11, gap=7)
bullet(tf, "Saved jobs & compare", "— shortlist and weigh two offers", size=11, gap=7)
bullet(tf, "Withdraw anytime", "— pull an application straight back from the tracker", size=11, gap=7)
bullet(tf, "Profile & photo", "— guided setup, parsed skills & avatar upload", size=11, gap=7)
bullet(tf, "Follow companies", "— one tap keeps their openings in view", size=11, gap=0)

rx = Inches(6.85)
b1 = browser_frame(s, C("s3-candidate.png"), rx, cy, cw2)
t = tb(s, rx, b1 + Inches(0.1), Inches(5.6), Inches(0.3))
p = para(t, first=True)
run(p, "Candidate dashboard — live stats, tracker & interviews", 10.5, MUTED, italic=True)

wy = b1 + Inches(0.52)
rect(s, rx, wy, cw2, Inches(1.6), PAPER, line=LINEC, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.09)
t = tb(s, rx + Inches(0.28), wy + Inches(0.16), cw2 - Inches(0.5), Inches(0.3))
p = para(t, first=True)
run(p, "ALSO BUILT IN", 11, INK, bold=True, spacing=2)
x = rx + Inches(0.28)
for txt in ("Withdraw application", "Follow companies"):
    x = chip(s, x, wy + Inches(0.56), txt, WHITE, BLUE, size=10)
x = rx + Inches(0.28)
for txt in ("Avatar upload", "Deadline alerts", "Compare jobs"):
    x = chip(s, x, wy + Inches(1.02), txt, WHITE, TEAL, size=10)

pagenum(s, 4)

# ============================================================
# SLIDE 5 — FUNCTIONALITIES · RECRUITERS
# ============================================================
s = slide()
rect(s, 0, 0, SW, SH, WHITE)
rect(s, 0, 0, SW, Inches(0.12), AMBER)
kicker(s, Inches(0.6), Inches(0.4), "04", "Project Functionalities · Recruiters", color=AMBER)
title(s, Inches(0.6), Inches(0.7), Inches(12.2), "For recruiters — post, screen, schedule and hire", size=26)

cy = Inches(1.5); ch2 = Inches(4.75); cx = Inches(0.6); cw2 = Inches(6.05)
rect(s, cx, cy, cw2, ch2, RGBColor(0xFD, 0xF3, 0xDF), line=RGBColor(0xF3, 0xDD, 0xAE), lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.055)
t = tb(s, cx + Inches(0.3), cy + Inches(0.2), cw2 - Inches(0.6), Inches(0.35))
p = para(t, first=True)
run(p, "FOR RECRUITERS", 12.5, AMBERD, bold=True, spacing=2)
tf = tb(s, cx + Inches(0.3), cy + Inches(0.66), cw2 - Inches(0.55), ch2 - Inches(0.8))
bullet(tf, "Post & manage jobs", "— skills, benefits, deadlines & status", mark_color=AMBER, first=True, size=11, gap=7)
bullet(tf, "Visual pipeline", "— stage moves with live candidate counts", mark_color=AMBER, size=11, gap=7)
bullet(tf, "Candidate screening", "— search, filters & full profiles", mark_color=AMBER, size=11, gap=7)
bullet(tf, "Saved leads", "— candidates who saved your job but never applied", mark_color=AMBER, size=11, gap=7)
bullet(tf, "Interview scheduling", "— date, validated meeting link & status history", mark_color=AMBER, size=11, gap=7)
bullet(tf, "Company branding", "— mission, gallery, benefits & reviews", mark_color=AMBER, size=11, gap=7)
bullet(tf, "Contact inbox", "— site enquiries stored, recruiters only", mark_color=AMBER, size=11, gap=7)
bullet(tf, "Live analytics & chat", "— charts plus real-time messaging", mark_color=AMBER, size=11, gap=0)

rx = Inches(6.85)
b2 = browser_frame(s, C("s3-recruiter.png"), rx, cy, cw2)
t = tb(s, rx, b2 + Inches(0.1), Inches(5.6), Inches(0.3))
p = para(t, first=True)
run(p, "Recruiter dashboard — pipeline, applicants & analytics", 10.5, MUTED, italic=True)

wy = b2 + Inches(0.52)
rect(s, rx, wy, cw2, Inches(1.6), PAPER, line=LINEC, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.09)
t = tb(s, rx + Inches(0.28), wy + Inches(0.16), cw2 - Inches(0.5), Inches(0.3))
p = para(t, first=True)
run(p, "RECRUITER TOOLKIT", 11, INK, bold=True, spacing=2)
x = rx + Inches(0.28)
for txt in ("Saved leads", "Contact inbox", "Close jobs"):
    x = chip(s, x, wy + Inches(0.56), txt, WHITE, AMBERD, size=10)
x = rx + Inches(0.28)
for txt in ("Delete reviews", "Status history", "Interview links"):
    x = chip(s, x, wy + Inches(1.02), txt, WHITE, TEAL, size=10)

pagenum(s, 5)

# ============================================================
# SLIDE 6 — PLATFORM & ENGAGEMENT LAYER
# ============================================================
s = slide()
rect(s, 0, 0, SW, SH, PAPER)
rect(s, 0, 0, Inches(0.14), SH, PURPLE)
kicker(s, Inches(0.6), Inches(0.42), "05", "Platform & Engagement Layer", color=PURPLE)
title(s, Inches(0.6), Inches(0.72), Inches(12.1), "Everything between \u201capply\u201d and \u201chire\u201d", size=26)
subtitle(s, Inches(0.6), Inches(1.2), Inches(12.1),
         "Realtime, social and messaging primitives that keep both sides coming back — all served by the same JWT-authenticated API.")

features = [
    ("Realtime chat", "1:1 messaging with read receipts, presence & history", BLUE),
    ("Typing indicators", "live \u201ctyping…\u201d state on both sides of the chat", TEAL),
    ("Notification bell", "unread counts, mark-one & mark-all-read", PURPLE),
    ("Follow companies", "persisted follower lists with live counts", GREEN),
    ("Withdraw application", "candidates exit the pipeline themselves", AMBER),
    ("Saved-lead mining", "warm intent from saved-but-not-applied candidates", RGBColor(0xE1, 0x1D, 0x48)),
    ("Contact inbox", "enquiries persisted in Mongo, gated to recruiters", TEAL),
    ("Paged & safe search", "server-side pagination, regex-escaped queries", DARK2),
]
fw2, fh2 = Inches(2.98), Inches(1.98)
gx, gy2 = Inches(0.6), Inches(1.92)
gapx, gapy = Inches(0.18), Inches(0.2)
for i, (h1, b, c) in enumerate(features):
    xx = gx + (i % 4) * (fw2 + gapx)
    yy = gy2 + (i // 4) * (fh2 + gapy)
    rect(s, xx, yy, fw2, fh2, WHITE, line=LINEC, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    rect(s, xx, yy, Inches(0.09), fh2, c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    t = tb(s, xx + Inches(0.26), yy + Inches(0.16), Inches(0.9), Inches(0.35))
    p = para(t, first=True)
    run(p, f"{i+1:02d}", 16, c, bold=True)
    t = tb(s, xx + Inches(0.26), yy + Inches(0.58), fw2 - Inches(0.45), Inches(0.5))
    p = para(t, first=True, line=1.05)
    run(p, h1, 12.5, INK, bold=True)
    t = tb(s, xx + Inches(0.26), yy + Inches(1.1), fw2 - Inches(0.45), Inches(0.78))
    p = para(t, first=True, line=1.12)
    run(p, b, 9.8, MUTED)

sy4 = Inches(6.24)
rect(s, Inches(0.6), sy4, Inches(12.13), Inches(0.62), DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.32)
t = tb(s, Inches(0.9), sy4, Inches(11.53), Inches(0.62), anchor=MSO_ANCHOR.MIDDLE)
p = para(t, first=True, align=PP_ALIGN.CENTER)
run(p, "Always on — deadline enforcement · duplicate-apply blocking · expired listings hidden · role-gated endpoints", 10.5, WHITE, bold=True)

pagenum(s, 6)

# ============================================================
# SLIDE 7 — METHODOLOGY / ARCHITECTURE
# ============================================================
s = slide()
rect(s, 0, 0, SW, SH, PAPER)
rect(s, SW - Inches(0.14), 0, Inches(0.14), SH, BLUE)
kicker(s, Inches(0.6), Inches(0.42), "06", "Methodology / Architecture")
title(s, Inches(0.6), Inches(0.72), Inches(9.5), "Three-tier architecture, secured by design")

layers = [
    ("FRONTEND", "HTML5 · CSS3 · Vanilla JS · Chart.js · Socket.io · Font Awesome 6 · DM Sans · Dark Mode", BLUE),
    ("BACKEND",  "Node.js · Express.js 5 · Socket.io · JWT · bcryptjs · multer · pdfjs-dist", TEAL),
    ("DATABASE", "MongoDB Atlas · Mongoose ODM · 8 collections · aggregation pipelines", GREEN),
]
for (name, techs, c), yy in zip(layers, (Inches(1.62 + i * 1.12) for i in range(3))):
    rect(s, Inches(0.6), yy, Inches(7.15), Inches(0.92), WHITE, line=LINEC, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    rect(s, Inches(0.6), yy, Inches(0.1), Inches(0.92), c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    t = tb(s, Inches(0.92), yy + Inches(0.09), Inches(6.6), Inches(0.32))
    p = para(t, first=True)
    run(p, name, 12.5, c, bold=True, spacing=2)
    t = tb(s, Inches(0.92), yy + Inches(0.42), Inches(6.65), Inches(0.45))
    p = para(t, first=True)
    run(p, techs, 10.2, MUTED)

# methodology stepper
sy = Inches(4.98)
t = tb(s, Inches(0.6), sy, Inches(2.3), Inches(0.4), anchor=MSO_ANCHOR.MIDDLE)
p = para(t, first=True)
run(p, "METHODOLOGY", 11, INK, bold=True, spacing=1)
x = Inches(2.55)
for txt, c in (("Requirement Analysis", BLUE), ("UI/UX Design", PURPLE),
               ("API Development", TEAL), ("Integration & Testing", AMBER), ("Deployment", GREEN)):
    x = chip(s, x, sy - Inches(0.05), txt, WHITE, c, size=10)
    if txt != "Deployment":
        arrow(s, x - Inches(0.01), sy + Inches(0.12), Inches(0.16), LINEC, h=0.02)

# security strip
sy2 = Inches(5.42)
t = tb(s, Inches(0.6), sy2, Inches(2.3), Inches(0.4), anchor=MSO_ANCHOR.MIDDLE)
p = para(t, first=True)
run(p, "SECURITY LAYER", 11, INK, bold=True, spacing=1)
x = Inches(2.55)
for txt in ("bcrypt hashing", "JWT + socket auth", "role-based access", "output encoding", "rate limiting"):
    x = chip(s, x, sy2 - Inches(0.05), txt, SKY, RGBColor(0x1D, 0x4E, 0xD8), size=10)

# bottom: full-width request-flow block diagram
rect(s, Inches(0.6), Inches(5.85), Inches(12.13), Inches(1.2), WHITE, line=LINEC, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
t = tb(s, Inches(0.88), Inches(5.97), Inches(3), Inches(0.24))
p = para(t, first=True)
run(p, "REQUEST FLOW", 10.5, INK, bold=True, spacing=2)
t = tb(s, Inches(4.3), Inches(5.99), Inches(8.15), Inches(0.24))
p = para(t, first=True, align=PP_ALIGN.RIGHT)
run(p, "Auth: login → bcrypt verify → signed JWT → role-gated REST + authenticated sockets", 8.5, MUTED, italic=True)

flow = [
    ("Client", "browser · fetch API", DARK2, RGBColor(0xD7, 0xE4, 0xFF)),
    ("Express.js API", "49 endpoints · middleware", BLUE, RGBColor(0xD7, 0xE4, 0xFF)),
    ("Auth & realtime", "JWT · Socket.io", PURPLE, RGBColor(0xE4, 0xDC, 0xFC)),
    ("Mongoose ODM", "models & schemas", TEAL, RGBColor(0xC9, 0xE8, 0xF3)),
    ("MongoDB Atlas", "8 collections", GREEN, RGBColor(0xC7, 0xF2, 0xE4)),
]
fw, fh = Inches(2.12), Inches(0.6)
for i, (h1, sub, fill, scol) in enumerate(flow):
    xx = Inches(0.88) + i * (fw + Inches(0.28))
    rect(s, xx, Inches(6.27), fw, fh, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16)
    t = tb(s, xx, Inches(6.34), fw, Inches(0.26))
    p = para(t, first=True, align=PP_ALIGN.CENTER)
    run(p, h1, 11.5, WHITE, bold=True)
    t = tb(s, xx, Inches(6.58), fw, Inches(0.24))
    p = para(t, first=True, align=PP_ALIGN.CENTER)
    run(p, sub, 8.5, scol)
    if i < 4:
        arrow(s, xx + fw + Inches(0.04), Inches(6.57), Inches(0.2), LBLUE)

pagenum(s, 7)

# ============================================================
# SLIDE 8 — SECURITY & DATA INTEGRITY
# ============================================================
s = slide()
rect(s, 0, 0, SW, SH, WHITE)
rect(s, 0, 0, SW, Inches(0.12), PURPLE)
kicker(s, Inches(0.6), Inches(0.4), "07", "Security & Data Integrity", color=PURPLE)
title(s, Inches(0.6), Inches(0.7), Inches(12.2), "Secured at every layer — not just at login", size=26)

spy, sph = Inches(1.5), Inches(3.9)
lpw = Inches(6.05)
rect(s, Inches(0.6), spy, lpw, sph, SKY, line=RGBColor(0xC7, 0xDA, 0xF8), lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.055)
t = tb(s, Inches(0.9), spy + Inches(0.2), lpw - Inches(0.6), Inches(0.35))
p = para(t, first=True)
run(p, "AUTHENTICATION & ACCESS", 12.5, BLUE, bold=True, spacing=2)
tf = tb(s, Inches(0.9), spy + Inches(0.66), lpw - Inches(0.55), sph - Inches(0.8))
bullet(tf, "Candidate-only registration", "— recruiter accounts provisioned by the administrator", first=True, size=11, gap=9)
bullet(tf, "JWT everywhere", "— REST headers plus the socket handshake; invalid tokens rejected at connect", size=11, gap=9)
bullet(tf, "Role gates per route", "— candidate-only and recruiter-only endpoints enforced", size=11, gap=9)
bullet(tf, "Ownership checks", "— conversations, applications, inbox & reviews verified per user", size=11, gap=9)
bullet(tf, "Rate-limited auth", "— login and registration throttled against brute force", size=11, gap=8)
bullet(tf, "Bcrypt hashing", "— credentials stored as salted hashes, never plaintext", size=11, gap=0)

rpx = Inches(6.85)
rect(s, rpx, spy, lpw, sph, RGBColor(0xF1, 0xEC, 0xFE), line=RGBColor(0xDD, 0xD3, 0xFB), lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.055)
t = tb(s, rpx + Inches(0.3), spy + Inches(0.2), lpw - Inches(0.6), Inches(0.35))
p = para(t, first=True)
run(p, "VALIDATION & DATA INTEGRITY", 12.5, PURPLE, bold=True, spacing=2)
tf = tb(s, rpx + Inches(0.3), spy + Inches(0.66), lpw - Inches(0.55), sph - Inches(0.8))
bullet(tf, "Output encoding", "— user text escaped before it reaches the DOM (XSS)", mark_color=PURPLE, first=True, size=11, gap=9)
bullet(tf, "Safe links", "— interview URLs restricted to http:// and https://", mark_color=PURPLE, size=11, gap=9)
bullet(tf, "Review integrity", "— 1–5 rating validation, one review per user, recomputed aggregates", mark_color=PURPLE, size=11, gap=9)
bullet(tf, "Deadline enforcement", "— expired postings hidden, late applications rejected", mark_color=PURPLE, size=11, gap=9)
bullet(tf, "Input hygiene", "— regex-escaped search, ObjectId guards, length limits", mark_color=PURPLE, size=11, gap=8)
bullet(tf, "Duplicate guards", "— one application per job, one review per company", mark_color=PURPLE, size=11, gap=0)

vy = Inches(5.72)
rect(s, Inches(0.6), vy, Inches(12.13), Inches(0.86), DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
t = tb(s, Inches(0.95), vy + Inches(0.13), Inches(11.4), Inches(0.26))
p = para(t, first=True)
run(p, "VERIFIED IN TESTS", 10.5, AMBER, bold=True, spacing=2)
x = Inches(0.95)
for txt in ("anonymous sockets rejected", "cross-role probes → 401 / 403", "duplicate application blocked", "XSS payloads rendered as text"):
    x = chip(s, x, vy + Inches(0.44), txt, DARK3, RGBColor(0xBF, 0xD3, 0xF5), size=9.5)

pagenum(s, 8)

# ============================================================
# SLIDE 9 — RESULTS
# ============================================================
s = slide()
rect(s, 0, 0, SW, SH, WHITE)
rect(s, 0, 0, SW, Inches(0.12), AMBER)
kicker(s, Inches(0.6), Inches(0.42), "08", "Results", color=AMBER)
title(s, Inches(0.6), Inches(0.72), Inches(9.5), "Outputs achieved — a working platform")

kpis = [
    ("49", "REST endpoints", BLUE),
    ("8", "data collections", GREEN),
    ("20k+", "lines of code", TEAL),
    ("12", "responsive pages", AMBER),
]
kw, kh = Inches(2.72), Inches(1.32)
for i, (num, lab, c) in enumerate(kpis):
    xx = Inches(0.6) + i * (kw + Inches(0.24))
    rect(s, xx, Inches(1.6), kw, kh, PAPER, line=LINEC, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.09)
    t = tb(s, xx, Inches(1.72), kw, Inches(0.6))
    p = para(t, first=True, align=PP_ALIGN.CENTER)
    run(p, num, 29, c, bold=True)
    t = tb(s, xx, Inches(2.42), kw, Inches(0.4))
    p = para(t, first=True, align=PP_ALIGN.CENTER)
    run(p, lab, 11, MUTED, bold=True)

t = tb(s, Inches(0.6), Inches(3.28), Inches(6), Inches(0.4))
p = para(t, first=True)
run(p, "DEMONSTRATED OUTPUTS", 12, INK, bold=True, spacing=2)
tf = tb(s, Inches(0.6), Inches(3.68), Inches(6.15), Inches(2.8))
bullet(tf, "End-to-end hiring loop", "— post, apply, screen, schedule & hire", first=True)
bullet(tf, "ATS resume scanner", "— PDF parsing, resume score & job-fit ranking")
bullet(tf, "Real-time collaboration", "— chat, typing, presence & read receipts")
bullet(tf, "Rich company pages", "— 7 companies, culture pages & verified reviews")
bullet(tf, "Engagement engine", "— follows, withdrawals & warm saved leads")
bullet(tf, "Analytics dashboards", "— Chart.js visuals for both roles")

bot = browser_frame(s, C("s6-ats.png"), Inches(7.15), Inches(3.32), Inches(5.55))
t = tb(s, Inches(7.15), bot + Inches(0.12), Inches(5.5), Inches(0.3))
p = para(t, first=True)
run(p, "ATS scanner — resume score, fixes & live job-fit ranking", 10.5, MUTED, italic=True)

ry = Inches(6.72)
t = tb(s, Inches(0.6), ry - Inches(0.02), Inches(1.3), Inches(0.4), anchor=MSO_ANCHOR.MIDDLE)
p = para(t, first=True)
run(p, "ALSO VERIFIED", 11, INK, bold=True, spacing=1)
x = Inches(2.0)
for txt, c in (("Authenticated realtime", BLUE), ("Mobile-first layout", GREEN), ("Zero build step", TEAL)):
    x = chip(s, x, ry - Inches(0.08), txt, PAPER, c, size=10)

pagenum(s, 9)

# ============================================================
# SLIDE 10 — CONCLUSION & FUTURE WORK
# ============================================================
s = slide()
rect(s, 0, 0, SW, SH, DARK)
circle(s, SW - Inches(1.2), Inches(-0.6), Inches(2.4), DARK3, alpha=55)
circle(s, Inches(-0.8), SH + Inches(0.5), Inches(2.1), BLUE, alpha=16)
kicker(s, Inches(0.6), Inches(0.45), "09", "Conclusion & Future Work", on_dark=True)
title(s, Inches(0.6), Inches(0.75), Inches(11), "Where the project stands — and where it goes next", color=WHITE, size=28)

# conclusion panel
cx2, cy2, cw3, ch3 = Inches(0.6), Inches(1.62), Inches(6.05), Inches(4.15)
rect(s, cx2, cy2, cw3, ch3, DARK2, line=RGBColor(0x2E, 0x4C, 0x85), lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.055)
t = tb(s, cx2 + Inches(0.3), cy2 + Inches(0.2), cw3 - Inches(0.6), Inches(0.35))
p = para(t, first=True)
run(p, "CONCLUSION", 12.5, LBLUE, bold=True, spacing=2)
tf = tb(s, cx2 + Inches(0.3), cy2 + Inches(0.62), cw3 - Inches(0.55), ch3 - Inches(0.8))
bullet(tf, "Fully functional platform", "— every core flow works end-to-end on live data", dark=True, first=True, size=11, gap=7)
bullet(tf, "Both journeys served", "— candidates track applications; recruiters manage a pipeline", dark=True, size=11, gap=7)
bullet(tf, "Resume intelligence", "— PDF parsing, ATS scoring & job-fit matching", dark=True, size=11, gap=7)
bullet(tf, "Real-time collaboration", "— chat, typing, presence, read receipts & alerts", dark=True, size=11, gap=7)
bullet(tf, "Security-first stack", "— JWT + socket auth, ownership checks & rate limiting", dark=True, size=11, gap=7)
bullet(tf, "Lean by design", "— framework-free frontend: zero build step, instant deploys", dark=True, size=11, gap=0)

# future work panel
fx2 = Inches(6.85)
rect(s, fx2, cy2, cw3, ch3, DARK3, line=RGBColor(0x3A, 0x5B, 0x99), lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.055)
t = tb(s, fx2 + Inches(0.3), cy2 + Inches(0.2), cw3 - Inches(0.6), Inches(0.35))
p = para(t, first=True)
run(p, "FUTURE ENHANCEMENTS", 12.5, AMBER, bold=True, spacing=2)
future = [
    ("Email & push alerts", "notify candidates outside the in-app bell", BLUE),
    ("AI-assisted screening", "LLM ranking beyond rule-based ATS scores", PURPLE),
    ("Admin verification", "company and posting approval workflow", TEAL),
    ("Multi-recruiter teams", "per-company workspaces, roles & permissions", AMBER),
    ("PWA mobile app", "installable experience with offline support", GREEN),
]
fy = cy2 + Inches(0.66)
for i, (h1, b, c) in enumerate(future):
    yy = fy + i * Inches(0.66)
    circle(s, fx2 + Inches(0.5), yy + Inches(0.19), Inches(0.16), c)
    t = tb(s, fx2 + Inches(0.78), yy + Inches(0.02), cw3 - Inches(1.0), Inches(0.3))
    p = para(t, first=True)
    run(p, h1, 12, WHITE, bold=True)
    t = tb(s, fx2 + Inches(0.78), yy + Inches(0.28), cw3 - Inches(1.0), Inches(0.3))
    p = para(t, first=True)
    run(p, b, 9.5, RGBColor(0xAD, 0xC2, 0xE4))

t = tb(s, Inches(1.5), Inches(6.25), SW - Inches(3.0), Inches(0.45))
p = para(t, first=True, align=PP_ALIGN.CENTER)
run(p, "Thank you", 17, WHITE, bold=True)
run(p, "  ·  Questions welcome", 14, RGBColor(0x9F, 0xB4, 0xD8), italic=True)

brandfoot(s, on_dark=True)
pagenum(s, 10, on_dark=True)

prs.save(OUT)
print("Saved:", OUT)
